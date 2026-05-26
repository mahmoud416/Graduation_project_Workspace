"""Quality Control (QC) API routes."""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from pathlib import Path
from typing import Annotated, Any, Dict, List, Optional
from uuid import uuid4

from fastapi import (
    APIRouter,
    Depends,
    File,
    Form,
    HTTPException,
    Query,
    UploadFile,
    status,
)
from fastapi.responses import StreamingResponse

from app.db.mongodb import get_database
from app.dependencies.auth import get_current_user
from app.dependencies.rbac import require_quality_control_user, ensure_roles
from app.schemas.quality_control import (
    DatasetReference,
    QualityStandardCreate,
    QualityStandardUpdate,
    QualityAnalysisRequest,
)
from app.services.quality_standard_service import QualityStandardService
from app.services.quality_analysis_service import QualityAnalysisService
from app.services.quality_report_service import QualityReportService

try:  # pragma: no cover - optional dependency
    from fpdf import FPDF
except ImportError:  # pragma: no cover - handled at runtime
    FPDF = None  # type: ignore

try:  # pragma: no cover - optional dependency
    import pdfplumber
except ImportError:  # pragma: no cover - handled at runtime
    pdfplumber = None  # type: ignore

try:  # pragma: no cover - optional dependency
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover - handled at runtime
    DocxDocument = None  # type: ignore


router = APIRouter(prefix="/qc", tags=["Quality Control"])

UPLOADS_ROOT = Path(__file__).resolve().parents[2] / "uploads"
QC_DATASET_ROOT = UPLOADS_ROOT / "qc_datasets"
QC_DATASET_ROOT.mkdir(parents=True, exist_ok=True)


# ---------------------------------------------------------------------------
# Quality Standards
# ---------------------------------------------------------------------------


@router.post("/standards", status_code=status.HTTP_201_CREATED)
async def create_standard(
    payload: QualityStandardCreate,
    current_user=Depends(require_quality_control_user),
    db=Depends(get_database),
):
    doc = await QualityStandardService.create_standard(
        db,
        payload,
        created_by=current_user["_id"],
    )
    return doc


@router.get("/standards")
async def list_standards(
    project_id: Optional[str] = Query(None),
    status_filter: Optional[str] = Query("active"),
    current_user=Depends(get_current_user),
    db=Depends(get_database),
):
    ensure_roles(current_user, ["quality_control", "admin", "sub_admin", "manager"])
    normalized_status = status_filter if status_filter not in (None, "", "all") else None
    docs = await QualityStandardService.list_standards(
        db,
        project_id=project_id,
        status=normalized_status,
    )
    return docs


@router.get("/standards/{standard_id}")
async def get_standard(
    standard_id: str,
    current_user=Depends(require_quality_control_user),
    db=Depends(get_database),
):
    try:
        return await QualityStandardService.get_standard(db, standard_id)
    except ValueError as exc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail=str(exc)) from exc


@router.put("/standards/{standard_id}")
async def update_standard(
    standard_id: str,
    payload: QualityStandardUpdate,
    current_user=Depends(require_quality_control_user),
    db=Depends(get_database),
):
    try:
        return await QualityStandardService.update_standard(db, standard_id, payload)
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


@router.delete("/standards/{standard_id}", status_code=status.HTTP_204_NO_CONTENT)
async def archive_standard(
    standard_id: str,
    current_user=Depends(require_quality_control_user),
    db=Depends(get_database),
):
    try:
        await QualityStandardService.archive_standard(db, standard_id)
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
    return None


@router.post("/standards/{standard_id}/dataset", status_code=status.HTTP_201_CREATED)
async def upload_dataset_reference(
    standard_id: str,
    file: UploadFile = File(...),
    tags: Optional[str] = Form(None),
    file_type: str = Form("image"),
    current_user=Depends(require_quality_control_user),
    db=Depends(get_database),
):
    saved = await _store_dataset_file(standard_id, file)
    tag_list = _parse_tags(tags)
    dataset_ref = DatasetReference(
        file_id=uuid4().hex,
        file_path=saved["relative_path"],
        file_type=file_type or "image",
        tags=tag_list,
        size_bytes=saved["size"],
    )
    try:
        doc = await QualityStandardService.add_dataset_reference(db, standard_id, dataset_ref)
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
    return doc


# ---------------------------------------------------------------------------
# AI Analysis
# ---------------------------------------------------------------------------


@router.post("/tasks/{task_id}/analyze", status_code=status.HTTP_202_ACCEPTED)
async def analyze_task_with_ai(
    task_id: str,
    project_id: str = Form(...),
    task_title: str = Form(...),
    task_description: str = Form(""),
    description_override: Optional[str] = Form(None),
    standard_ids: Optional[str] = Form(None, description="Comma-separated or JSON array of standard IDs"),
    report_type: Optional[str] = Form(None, description="Report type key from REPORT_TYPES"),
    document_files: Annotated[Optional[List[UploadFile]], File(description="Optional document uploads")]
    = None,
    image_files: Annotated[Optional[List[UploadFile]], File(description="Optional image uploads")]
    = None,
    current_user=Depends(get_current_user),
    db=Depends(get_database),
):
    from app.services import qc_service as _qc_svc
    if report_type and report_type not in _qc_svc.REPORT_TYPES:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            detail="Unknown report type. Use GET /quality/report-types for valid keys.",
        )

    parsed_standard_ids = _parse_standard_ids(standard_ids)
    text_payloads = await _read_documents(document_files)
    image_payloads = await _read_images(image_files)

    analysis_request = QualityAnalysisRequest(
        task_id=task_id,
        project_id=project_id,
        task_title=task_title,
        task_description=task_description,
        description_override=description_override,
        standard_ids=parsed_standard_ids,
        report_type=report_type,
    )

    try:
        analysis = await QualityAnalysisService.run_analysis(
            db,
            analysis_request,
            triggered_by=current_user["_id"],
            file_texts=text_payloads,
            image_bytes=image_payloads,
        )
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc

    return analysis


@router.get("/tasks/{task_id}/analysis/latest")
async def latest_task_analysis(
    task_id: str,
    current_user=Depends(get_current_user),
    db=Depends(get_database),
):
    doc = await QualityAnalysisService.latest_for_task(db, task_id)
    if not doc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail="No analysis found for this task")
    return doc


@router.get("/analyses")
async def list_analyses(
    project_id: Optional[str] = Query(None),
    skip: int = Query(0, ge=0),
    limit: int = Query(20, ge=1, le=100),
    current_user=Depends(require_quality_control_user),
    db=Depends(get_database),
):
    docs = await QualityAnalysisService.list_analyses(
        db,
        project_id=project_id,
        limit=limit,
        skip=skip,
    )
    return docs


@router.get("/analyses/{analysis_id}")
async def get_analysis(
    analysis_id: str,
    current_user=Depends(require_quality_control_user),
    db=Depends(get_database),
):
    try:
        return await QualityAnalysisService.get_analysis(db, analysis_id)
    except ValueError as exc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, detail=str(exc)) from exc


# ---------------------------------------------------------------------------
# Reports & Exports
# ---------------------------------------------------------------------------


@router.get("/reports/overview")
async def reports_overview(
    project_id: Optional[str] = Query(None),
    days: int = Query(30, ge=1, le=180),
    current_user=Depends(require_quality_control_user),
    db=Depends(get_database),
):
    return await QualityReportService.quality_overview(db, project_id=project_id, days=days)


@router.get("/reports/export")
async def export_reports(
    format: str = Query("csv", pattern="^(csv|pdf)$"),
    project_id: Optional[str] = Query(None),
    days: int = Query(30, ge=1, le=365),
    current_user=Depends(require_quality_control_user),
    db=Depends(get_database),
):
    data = await QualityReportService.quality_overview(db, project_id=project_id, days=days)
    filename_suffix = project_id or "all-projects"
    if format == "csv":
        csv_buffer = _build_csv_export(data)
        headers = {"Content-Disposition": f"attachment; filename=qc-report-{filename_suffix}.csv"}
        return StreamingResponse(
            iter([csv_buffer.getvalue().encode("utf-8")]),
            media_type="text/csv",
            headers=headers,
        )

    if format == "pdf":
        if FPDF is None:
            raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, detail="PDF export requires fpdf2 dependency")
        pdf_bytes = _build_pdf_export(data)
        headers = {"Content-Disposition": f"attachment; filename=qc-report-{filename_suffix}.pdf"}
        return StreamingResponse(io.BytesIO(pdf_bytes), media_type="application/pdf", headers=headers)

    raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Unsupported export format")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


async def _store_dataset_file(standard_id: str, upload: UploadFile) -> dict[str, str | int]:
    if not upload.filename:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Missing filename for dataset upload")
    contents = await upload.read()
    if not contents:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, detail="Uploaded file is empty")
    safe_standard = standard_id.replace("/", "_")
    destination_dir = QC_DATASET_ROOT / safe_standard
    destination_dir.mkdir(parents=True, exist_ok=True)
    unique_name = f"{uuid4().hex}_{Path(upload.filename).name}"
    destination = destination_dir / unique_name
    destination.write_bytes(contents)
    relative_path = destination.relative_to(UPLOADS_ROOT).as_posix()
    return {"relative_path": relative_path, "size": len(contents)}


def _parse_tags(raw: Optional[str]) -> List[str]:
    if not raw:
        return []
    try:
        if raw.strip().startswith("["):
            parsed = json.loads(raw)
            if isinstance(parsed, list):
                return [str(tag).strip() for tag in parsed if str(tag).strip()]
    except json.JSONDecodeError:
        pass
    return [tag.strip() for tag in raw.split(",") if tag.strip()]


def _parse_standard_ids(raw: Optional[str]) -> List[str]:
    if not raw:
        return []
    try:
        if raw.strip().startswith("["):
            parsed = json.loads(raw)
            if isinstance(parsed, list):
                return [str(item) for item in parsed if item]
    except json.JSONDecodeError:
        pass
    return [item.strip() for item in raw.split(",") if item.strip()]


async def _read_documents(files: Optional[List[UploadFile]]) -> List[dict]:
    uploads = _coerce_uploads(files)
    texts: List[dict] = []
    for upload in uploads:
        content_bytes = await upload.read()
        if not content_bytes:
            continue

        file_name = upload.filename or "document"
        content_type = upload.content_type or ""
        fname_lower = file_name.lower()
        is_pdf = fname_lower.endswith(".pdf") or "pdf" in content_type.lower()
        is_docx = fname_lower.endswith(".docx") or fname_lower.endswith(".doc") or "wordprocessingml" in content_type.lower()

        if is_pdf:
            text = _extract_pdf_text(content_bytes)
            file_type = "pdf"
        elif is_docx:
            text = _extract_docx_text(content_bytes)
            file_type = "docx"
        else:
            try:
                text = content_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text = content_bytes.decode("latin-1", errors="ignore")
            file_type = content_type or "document"

        import base64 as _b64
        entry: Dict[str, Any] = {
            "file_name": file_name,
            "content": text,
            "file_type": file_type,
        }
        if is_pdf:
            entry["raw_b64"] = _b64.b64encode(content_bytes).decode("utf-8")
        texts.append(entry)
    return texts


def _extract_pdf_text(content_bytes: bytes) -> str:
    if pdfplumber is None:
        return "[PDF extraction unavailable: install pdfplumber]"
    try:
        import io as _io
        pages_text = []
        with pdfplumber.open(_io.BytesIO(content_bytes)) as pdf:
            for i, page in enumerate(pdf.pages):
                page_text = page.extract_text() or ""
                # Also extract tables as text
                tables = page.extract_tables() or []
                table_text = ""
                for table in tables:
                    for row in table:
                        if row:
                            table_text += " | ".join(str(cell or "") for cell in row) + "\n"
                combined = page_text
                if table_text:
                    combined += f"\n[TABLE DATA]\n{table_text}"
                if combined.strip():
                    pages_text.append(f"[Page {i+1}]\n{combined.strip()}")
        return "\n\n".join(pages_text) if pages_text else "[PDF has no extractable text]"
    except Exception as e:
        return f"[PDF extraction error: {e}]"


def _extract_docx_text(content_bytes: bytes) -> str:
    if DocxDocument is None:
        return "[Word extraction unavailable: install python-docx]"
    try:
        import io as _io
        doc = DocxDocument(_io.BytesIO(content_bytes))
        parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(f"[TABLE ROW] {row_text}")
        return "\n".join(parts) if parts else "[Word document has no extractable text]"
    except Exception as e:
        return f"[Word extraction error: {e}]"


async def _read_images(files: Optional[List[UploadFile]]) -> List[bytes]:
    uploads = _coerce_uploads(files)
    images: List[bytes] = []
    for upload in uploads:
        content = await upload.read()
        if content:
            images.append(content)
    return images


def _coerce_uploads(value: Optional[List[UploadFile]]) -> List[UploadFile]:
    if not value:
        return []
    return [upload for upload in value if upload]


def _build_csv_export(data: Dict[str, Any]) -> io.StringIO:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["Date", "Average Score", "Evaluations"])
    for entry in data.get("scoreTrend", []):
        writer.writerow([entry.get("date"), entry.get("avgScore"), entry.get("evaluations")])
    writer.writerow([])
    writer.writerow(["Project", "Average Score", "Evaluations"])
    for project in data.get("projectScores", []):
        writer.writerow([project.get("projectId"), project.get("avgScore"), project.get("evaluations")])
    buffer.seek(0)
    return buffer


def _build_pdf_export(data: Dict[str, Any]) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=14)
    pdf.cell(0, 10, "Quality Control Report", ln=True)
    pdf.set_font("Helvetica", size=10)
    pdf.multi_cell(0, 8, f"Generated on: {datetime.utcnow().isoformat()} UTC")

    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Score Trend", ln=True)
    pdf.set_font("Helvetica", size=10)
    for entry in data.get("scoreTrend", [])[:20]:
        pdf.cell(0, 6, f"{entry.get('date')}: {entry.get('avgScore')} ({entry.get('evaluations')} evals)", ln=True)

    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Project Scores", ln=True)
    pdf.set_font("Helvetica", size=10)
    for project in data.get("projectScores", [])[:20]:
        pdf.cell(0, 6, f"{project.get('projectId')}: {project.get('avgScore')} ({project.get('evaluations')} evals)", ln=True)

    return pdf.output(dest="S").encode("latin-1")
